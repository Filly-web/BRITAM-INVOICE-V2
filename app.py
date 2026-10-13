from flask import Flask, render_template, request, send_file
from docx import Document
from docx.shared import Inches
from datetime import datetime
import sqlite3,qrcode,os
app=Flask(__name__)
UPLOAD_FOLDER='uploads'; OUTPUT_FOLDER='outputs'; QRCODE_FOLDER='qrcodes'; DATABASE='portal.db'
for f in [UPLOAD_FOLDER,OUTPUT_FOLDER,QRCODE_FOLDER]: os.makedirs(f,exist_ok=True)

def init_db():
    c=sqlite3.connect(DATABASE)
    c.execute('CREATE TABLE IF NOT EXISTS logs(id INTEGER PRIMARY KEY AUTOINCREMENT, filename TEXT, verification_code TEXT, verification_link TEXT, created_at TEXT)')
    c.commit(); c.close()
init_db()

@app.route('/')
def home():
    c=sqlite3.connect(DATABASE)
    total=c.execute('SELECT COUNT(*) FROM logs').fetchone()[0]
    today=c.execute("SELECT COUNT(*) FROM logs WHERE date(created_at)=date('now')").fetchone()[0]
    recent=c.execute('SELECT * FROM logs ORDER BY id DESC LIMIT 5').fetchall(); c.close()
    return render_template('index.html',total=total,today=today,recent=recent)

@app.route('/process',methods=['POST'])
def process():
    try:
        invoice=request.files['invoice']
        code=request.form['verification_code'].strip()
        hhmmss=request.form['verification_time'].replace(':','').strip()
        if not invoice.filename.lower().endswith('.docx'): return 'Only DOCX files are supported.'
        src=os.path.join(UPLOAD_FOLDER,invoice.filename); invoice.save(src)
        link=f'https://verify.tra.go.tz/{code}_{hhmmss}'
        qr=os.path.join(QRCODE_FOLDER,f'{code}.png'); qrcode.make(link).save(qr)
        doc=Document(src); inserted=False
        for p in doc.paragraphs:

    if '<<EFD_SECTION>>' in p.text:

        p.text = ''

        p.add_run('Verification Code: ' + code)

        p.add_run('\n')

        p.add_run(
            'Verification Link: ' + verification_link
        )

        p.add_run('\n')

        p.add_run().add_picture(
            qr_file,
            width=Inches(0.6)
        )

        inserted = True
        break
        if not inserted:
            doc.add_paragraph('Verification Code: '+code); doc.add_paragraph('Verification Link: '+link); doc.add_picture(qr,width=Inches(0.6))
        out=os.path.join(OUTPUT_FOLDER,'PROCESSED_'+invoice.filename); doc.save(out)
        c=sqlite3.connect(DATABASE)
        c.execute('INSERT INTO logs(filename,verification_code,verification_link,created_at) VALUES(?,?,?,?)',(invoice.filename,code,link,datetime.now().strftime('%Y-%m-%d %H:%M:%S')))
        c.commit(); c.close()
        return send_file(out,as_attachment=True)
    except Exception as e:
        return 'ERROR: '+str(e)

@app.route('/history')
def history():
    c=sqlite3.connect(DATABASE); rows=c.execute('SELECT * FROM logs ORDER BY id DESC').fetchall(); c.close()
    return render_template('history.html',records=rows)

if __name__=='__main__':
    port=int(os.environ.get('PORT',5000)); app.run(host='0.0.0.0',port=port)
